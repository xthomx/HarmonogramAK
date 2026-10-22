import io
import json
from datetime import datetime
from pathlib import Path
import tempfile
import types
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
import xlsx_format
from main import Api, atomic_write
import test_formats

class ExcelTests(unittest.TestCase):
    def setUp(self):
        fixture = test_formats.Formats(); fixture.setUp()
        self.plan = fixture.plan
        self.plan['calendarFilters']['specialization'] = ['spec:plc', '__none__']

    def edited(self, edit):
        workbook = load_workbook(io.BytesIO(xlsx_format.dumps(self.plan)))
        edit(workbook)
        out = io.BytesIO(); workbook.save(out); workbook.close()
        return out.getvalue()

    def test_complete_roundtrip(self):
        self.assertEqual(xlsx_format.loads(xlsx_format.dumps(self.plan)), self.plan)
        self.plan['people'] = []; self.plan['projects'] = []; self.plan['assignments'] = []
        self.assertEqual(xlsx_format.loads(xlsx_format.dumps(self.plan)), self.plan)

    def test_excel_types(self):
        w = load_workbook(io.BytesIO(xlsx_format.dumps(self.plan)))
        self.assertEqual(w.sheetnames, list(xlsx_format.SHEETS))
        self.assertEqual(w['Projekty']['D2'].value, '0012')
        self.assertEqual(w['Projekty']['B2'].data_type, 's')
        self.assertEqual(w['Projekty']['B2'].value, '=NieFormuła')
        self.assertIsInstance(w['Przypisania']['I2'].value, datetime)
        self.assertEqual(w['Przypisania']['K2'].value, 75)
        w.close()

    def test_excel_edits(self):
        def edit(w):
            w['Osoby']['C2'] = 'Elektryk'; w['Osoby']['D2'] = 'Nowa grupa'
            w['Przypisania']['I2'] = datetime(2026, 11, 1)
            w['Przypisania']['K2'] = 50
        plan = xlsx_format.loads(self.edited(edit))
        self.assertEqual(plan['people'][0]['role'], 'Elektryk')
        self.assertEqual(plan['people'][0]['group'], 'Nowa grupa')
        self.assertEqual(plan['assignments'][0]['start'], '2026-11-01')
        self.assertEqual(plan['assignments'][0]['load'], 50)

    def test_invalid_excel_rejected(self):
        for edit in [lambda w: w.remove(w['Osoby']),
                     lambda w: setattr(w['Osoby']['A1'], 'value', 'Wrong'),
                     lambda w: setattr(w['Osoby']['B2'], 'value', '=1+1'),
                     lambda w: setattr(w['Przypisania']['B2'], 'value', 'missing'),
                     lambda w: setattr(w['Przypisania']['K2'], 'value', 101),
                     lambda w: setattr(w['Przypisania']['I2'], 'value', '2027-12-31')]:
            with self.subTest(edit=edit), self.assertRaises(ValueError):
                xlsx_format.loads(self.edited(edit))
        with self.assertRaises(ValueError): xlsx_format.loads(b'not excel')

    def test_failed_save_preserves_workbook(self):
        with tempfile.TemporaryDirectory() as d:
            path = Path(d)/'plan.xlsx'
            atomic_write(path, json.dumps(self.plan)); before = path.read_bytes()
            with patch('main.os.replace', side_effect=PermissionError('locked')):
                with self.assertRaisesRegex(PermissionError, 'Excelu'):
                    atomic_write(path, json.dumps(self.plan))
            self.assertEqual(path.read_bytes(), before)
            self.assertEqual(list(Path(d).iterdir()), [path])

    def test_native_default_save_and_open(self):
        module = types.SimpleNamespace(FileDialog=types.SimpleNamespace(OPEN=1, SAVE=2))
        with tempfile.TemporaryDirectory() as d, patch.dict('sys.modules', {'webview': module}):
            path = Path(d)/'plan.xlsx'
            api = Api()
            api._window = types.SimpleNamespace(create_file_dialog=lambda *a, **k: [str(path)])
            self.assertEqual(api.save_plan(json.dumps(self.plan))['name'], 'plan.xlsx')
            result = api.open_plan()
            self.assertEqual(result['format'], 'xlsx')
            self.assertEqual(json.loads(result['text']), self.plan)
            api.accept_open(result['token'])
            self.assertEqual(api._path, path)
