import copy
import csv
import io
import unittest
from openpyxl import load_workbook
import fleet
import plan_format
import xlsx_format
from test_formats import Formats

class CarsTests(unittest.TestCase):
    def setUp(self):
        source=Formats();source.setUp();self.plan=source.plan
        self.plan['cars']=[{'id':'c1','name':'Skoda Octavia','plate':'SK 00123','seats':5,'usage':'Sekcyjne','owner':'PLC'},
                           {'id':'c2','name':'Auto 2','plate':'SK 00456','seats':2,'usage':'Indywidualne','owner':'Jan'}]
        self.plan['carReservations']=[{'id':'r1','car':'c1','project':'j1','start':'2026-11-01','end':'2026-11-05','passengers':4,'note':'=Tekst, nie formuła'}]

    def test_roundtrip_all_formats(self):
        for ext in ('.csv','.txt'):
            self.assertEqual(plan_format.loads(plan_format.dumps(self.plan,ext),ext),self.plan)
        self.assertEqual(xlsx_format.loads(xlsx_format.dumps(self.plan)),self.plan)

    def test_legacy_four_sheets(self):
        legacy=copy.deepcopy(self.plan);del legacy['cars'];del legacy['carReservations']
        book=load_workbook(io.BytesIO(xlsx_format.dumps(legacy)))
        book.remove(book['Samochody']);book.remove(book['Rezerwacje_aut'])
        out=io.BytesIO();book.save(out);book.close()
        self.assertEqual(xlsx_format.loads(out.getvalue()),legacy)

    def test_old_csv_headers_without_new_columns(self):
        legacy=copy.deepcopy(self.plan);del legacy['cars'];del legacy['carReservations']
        rows=list(csv.DictReader(io.StringIO(plan_format.dumps(legacy,'.csv')),delimiter=';'))
        fields=[f for f in plan_format.FIELDS if f not in plan_format.CAR_FIELDS]
        out=io.StringIO();writer=csv.DictWriter(out,fieldnames=fields,delimiter=';',extrasaction='ignore');writer.writeheader();writer.writerows(rows)
        self.assertEqual(plan_format.loads(out.getvalue(),'.csv'),legacy)

    def test_excel_numeric_cells_and_literal_notes(self):
        book=load_workbook(io.BytesIO(xlsx_format.dumps(self.plan)))
        self.assertEqual(book['Samochody']['D2'].value,5)
        self.assertEqual(book['Rezerwacje_aut']['I2'].value,4)
        self.assertEqual(book['Rezerwacje_aut']['J2'].data_type,'s')
        book.close()

    def test_validation_and_overlaps(self):
        for key,value in [('car','missing'),('project','missing'),('passengers',6),('start','2026-02-30'),('end','2025-01-01')]:
            p=copy.deepcopy(self.plan);p['carReservations'][0][key]=value
            with self.subTest(key=key),self.assertRaises(ValueError):fleet.validate(p)
        p=copy.deepcopy(self.plan);p['carReservations'].append(dict(p['carReservations'][0],id='r2',start='2026-11-05',end='2026-11-06'))
        with self.assertRaises(ValueError):fleet.validate(p)
        p['carReservations'][1]['start']='2026-11-06';fleet.validate(p)
        p=copy.deepcopy(self.plan);p['cars'][0]['seats']=3
        with self.assertRaises(ValueError):fleet.validate(p)

    def test_missing_new_sheet_rejected(self):
        book=load_workbook(io.BytesIO(xlsx_format.dumps(self.plan)));book.remove(book['Rezerwacje_aut'])
        out=io.BytesIO();book.save(out);book.close()
        with self.assertRaises(ValueError):xlsx_format.loads(out.getvalue())

    def test_expanded_car_fields_roundtrip_and_types(self):
        extra={'project':'j1','driver':'Jan Kowalski','route':'Polska / RO','odometer':123456,
               'vignetteRO':'2027-10-12','vignetteSK':'2027-01-01','vignetteHU':'2027-02-01',
               'vignetteCZ':'2027-03-01','vignetteAT':'2027-04-01','de':'Zielona naklejka',
               'inspection':'2027-10-12','tires':'Zimowe','kp':'O2 12.10.2026','notes':'=Nie formuła\nDruga linia'}
        self.plan['cars'][0].update(extra)
        for ext in ('.csv','.txt'):
            text=plan_format.dumps(self.plan,ext);self.assertIn('PLAN_ZESPOLU_3',text)
            self.assertEqual(plan_format.loads(text,ext),self.plan)
        raw=xlsx_format.dumps(self.plan);self.assertEqual(xlsx_format.loads(raw),self.plan)
        book=load_workbook(io.BytesIO(raw));sheet=book['Samochody'];headers={cell.value:cell.column for cell in sheet[1]}
        self.assertEqual(sheet.cell(2,headers['Stan_licznika']).value,123456)
        self.assertEqual(sheet.cell(2,headers['Uwagi_auta']).data_type,'s')
        self.assertEqual(sheet.cell(2,headers['Winieta_RO']).value.date().isoformat(),'2027-10-12');book.close()

    def test_previous_vehicle_sheet_without_new_columns(self):
        book=load_workbook(io.BytesIO(xlsx_format.dumps(self.plan)));sheet=book['Samochody']
        sheet.delete_cols(7,sheet.max_column-6)
        out=io.BytesIO();book.save(out);book.close()
        self.assertEqual(xlsx_format.loads(out.getvalue()),self.plan)

    def test_expanded_validation(self):
        for changes in ({'odometer':-1},{'odometer':1.5},{'vignetteSK':'2027-02-29'},{'project':'missing'}):
            p=copy.deepcopy(self.plan);p['cars'][0].update(changes)
            with self.subTest(changes=changes),self.assertRaises(ValueError):fleet.validate(p)
