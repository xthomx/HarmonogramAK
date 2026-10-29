"""Excel workbook interchange for the desktop planner."""
import csv
import io
from datetime import date, datetime
from zipfile import ZipFile, BadZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
import plan_format as table

SHEETS = {
    'Osoby': ('OSOBA', ['ID', 'Nazwa', 'Specjalizacja', 'Grupa']),
    'Projekty': ('PROJEKT', ['ID', 'Nazwa', 'Skrot', 'Nr_tematu', 'Opis', 'Kolor']),
    'Przypisania': ('PRZYPISANIE', ['ID', 'Osoba_ID', 'Projekt_ID', 'Osoba', 'Specjalizacja', 'Grupa', 'Projekt', 'Nr_tematu', 'Od', 'Do', 'Zaangazowanie_proc', 'Notatka', 'Etap']),
    'Samochody': ('SAMOCHOD', ['ID', 'Nazwa', 'Rejestracja', 'Miejsca', 'Uzytkowanie', 'Uzytkownik'] + list(table.CAR_EXTRA.values())),
    'Rezerwacje_aut': ('REZERWACJA_AUTA', ['ID', 'Samochod_ID', 'Nazwa', 'Rejestracja', 'Projekt_ID', 'Projekt', 'Od', 'Do', 'Liczba_osob', 'Notatka']),
    'Ustawienia': ('USTAWIENIE', ['Ustawienie', 'Wartosc']),
}


def dumps(plan):
    rows = [{k: table.unsafe(v) for k, v in r.items()}
            for r in csv.DictReader(io.StringIO(table.dumps(plan, '.csv')), delimiter=';')]
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, (kind, columns) in SHEETS.items():
        sheet = workbook.create_sheet(title)
        sheet.append(columns)
        records = [r for r in rows if r['Typ'] == kind]
        if title == 'Ustawienia':
            records.insert(0, {'Ustawienie': 'format', 'Wartosc': table.format_version(plan)})
        for record in records:
            row_number = sheet.max_row + 1
            for col_number, column in enumerate(columns, 1):
                value = record.get(column, '')
                cell = sheet.cell(row_number, col_number)
                if column in ('Od', 'Do','Winieta_RO','Winieta_SK','Winieta_HU','Winieta_CZ','Winieta_AT','Przeglad') and value:
                    cell.value = date.fromisoformat(value)
                    cell.number_format = 'yyyy-mm-dd'
                elif column in ('Zaangazowanie_proc','Miejsca','Liczba_osob','Stan_licznika') and value != '':
                    cell.value = int(value)
                else:
                    cell.value = value
                    cell.data_type = 's'  # Literal text, including leading zeros and '='.
                    cell.number_format = '@'
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='286553')
            sheet.column_dimensions[cell.column_letter].width = 24
    out = io.BytesIO()
    workbook.save(out)
    workbook.close()
    return out.getvalue()


def loads(raw):
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 50_000_000:
                raise ValueError('Rozpakowany plik Excel przekracza 50 MB.')
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, OSError, KeyError) as exc:
        raise ValueError('Nieprawidłowy plik Excel .xlsx.') from exc
    try:
        if set(workbook.sheetnames) not in ({'Osoby','Projekty','Przypisania','Ustawienia'}, set(SHEETS)):
            raise ValueError('Brak wymaganych arkuszy planu lub nieznany arkusz.')
        out = io.StringIO(newline='')
        writer = csv.DictWriter(out, fieldnames=table.FIELDS, delimiter=';')
        writer.writeheader()
        for title, (kind, columns) in SHEETS.items():
            if title not in workbook.sheetnames: continue
            sheet = workbook[title]
            if (sheet.max_row or 0) > 10002 or (sheet.max_column or 0) > len(columns):
                raise ValueError('Zbyt duży arkusz: ' + title)
            iterator = sheet.iter_rows()
            header = next(iterator, ())
            names = [cell.value for cell in header]
            expected = set(columns)
            allowed = [expected, expected - {'Etap'}] if kind == 'PRZYPISANIE' else [expected]
            if kind == 'SAMOCHOD': allowed.append(expected - set(table.CAR_EXTRA.values()))
            if len(names) != len(set(names)) or set(names) not in allowed:
                raise ValueError('Nieprawidłowe nagłówki w arkuszu: ' + title)
            for cells in iterator:
                if all(cell.value is None for cell in cells):
                    continue
                row = {'Typ': kind}
                for name, cell in zip(names, cells):
                    value = cell.value
                    if cell.data_type in ('f', 'e'):
                        raise ValueError(f'{title}!{cell.coordinate}: wpisz wartość zamiast formuły lub błędu.')
                    if isinstance(value, (datetime, date)):
                        if name not in ('Od', 'Do','Winieta_RO','Winieta_SK','Winieta_HU','Winieta_CZ','Winieta_AT','Przeglad'):
                            raise ValueError(f'{title}!{cell.coordinate}: oczekiwany tekst, nie data.')
                        value = value.date() if isinstance(value, datetime) else value
                        value = value.isoformat()
                    row[name] = '' if value is None else str(value)
                if kind == 'USTAWIENIE' and row['Ustawienie'] == 'format':
                    row['Typ'] = 'FORMAT'
                writer.writerow({k: table.safe(v) for k, v in row.items()})
        plan = table.loads(out.getvalue(), '.csv')
        if 'cars' in plan and not {'Samochody','Rezerwacje_aut'} <= set(workbook.sheetnames):
            raise ValueError('Brak arkuszy Samochody lub Rezerwacje_aut dla nowego formatu.')
        validate(plan)
        return plan
    finally:
        workbook.close()


def validate(plan):
    table.fleet.validate(plan)
    for key in ('people', 'projects', 'assignments'):
        ids = [item['id'] for item in plan[key]]
        if len(ids) > 10000 or any(not id.strip() for id in ids) or len(ids) != len(set(ids)):
            raise ValueError('Puste lub powtórzone ID albo zbyt wiele wierszy: ' + key)
    for item in plan['people'] + plan['projects']:
        if not item['name'].strip():
            raise ValueError('Nazwa osoby lub projektu nie może być pusta.')
    people = {p['id'] for p in plan['people']}
    projects = {p['id'] for p in plan['projects']}
    for item in plan['assignments']:
        if item['person'] not in people or item['project'] not in projects:
            raise ValueError('Przypisanie wskazuje nieistniejące Osoba_ID lub Projekt_ID.')
        try:
            start, end = (date.fromisoformat(item[k]) for k in ('start', 'end'))
        except ValueError as exc:
            raise ValueError('Daty przypisań muszą mieć format RRRR-MM-DD lub być datami Excela.') from exc
        if start.isoformat() != item['start'] or end.isoformat() != item['end'] or start > end:
            raise ValueError('Nieprawidłowy zakres dat przypisania.')
        if not 1 <= item['load'] <= 100:
            raise ValueError('Zaangażowanie musi być liczbą całkowitą od 1 do 100.')

